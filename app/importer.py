from __future__ import annotations

import json
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from .models import Formula, Measurement
from .storage import SQLStore


def _to_dict(v: Any) -> dict:
    if isinstance(v, dict):
        return v
    if v is None:
        return {}
    if isinstance(v, str):
        text = v.strip()
        if not text:
            return {}
        return json.loads(text)
    raise ValueError(f"value cannot convert to dict: {type(v)}")


def _to_bool(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return bool(v)
    if isinstance(v, str):
        text = v.strip().lower()
        return text in {"1", "true", "yes", "y", "是"}
    return False


def _headers(values: tuple[Any, ...]) -> list[str]:
    return [str(x).strip() if x is not None else "" for x in values]


def import_excel_dataset(content: bytes, store: SQLStore) -> dict:
    wb = load_workbook(BytesIO(content), data_only=True)
    report: dict[str, Any] = {
        "success_count": {"measurements": 0, "formulas": 0, "feedbacks": 0},
        "failed_rows": {"measurements": [], "formulas": [], "feedbacks": []},
    }

    sheet_aliases = {
        "measurements": ["measurements", "measurement", "测色", "测色数据"],
        "formulas": ["formulas", "formula", "配方", "配方库"],
        "feedbacks": ["feedbacks", "feedback", "反馈", "检测反馈"],
    }

    for dataset, aliases in sheet_aliases.items():
        ws = None
        for s in wb.worksheets:
            if s.title.strip().lower() in [x.lower() for x in aliases]:
                ws = s
                break
        if ws is None:
            continue

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        cols = _headers(rows[0])
        for idx, r in enumerate(rows[1:], start=2):
            if all(x is None or str(x).strip() == "" for x in r):
                continue
            row = {cols[i]: r[i] for i in range(min(len(cols), len(r)))}
            try:
                if dataset == "measurements":
                    m = Measurement(
                        id=str(row.get("id", "")).strip(),
                        car_model=str(row.get("car_model", "")).strip(),
                        color_code=str(row.get("color_code", "")).strip(),
                        angles=_to_dict(row.get("angles")),
                        texture=_to_dict(row.get("texture")),
                    )
                    if not m.id:
                        raise ValueError("measurement id is empty")
                    store.add_measurement(m)
                elif dataset == "formulas":
                    f = Formula(
                        id=str(row.get("id", "")).strip(),
                        name=str(row.get("name", "")).strip(),
                        pigments=_to_dict(row.get("pigments")),
                        linked_measurement_id=str(row.get("linked_measurement_id", "")).strip(),
                    )
                    if not f.id:
                        raise ValueError("formula id is empty")
                    if not store.has_measurement(f.linked_measurement_id):
                        raise ValueError(f"linked measurement not found: {f.linked_measurement_id}")
                    store.add_formula(f)
                else:
                    formula_id = str(row.get("formula_id", "")).strip()
                    if not formula_id:
                        raise ValueError("feedback formula_id is empty")
                    store.add_feedback(
                        formula_id=formula_id,
                        measured_performance=_to_dict(row.get("measured_performance")),
                        accepted=_to_bool(row.get("accepted")),
                        comment=str(row.get("comment", "") or ""),
                    )
                report["success_count"][dataset] += 1
            except Exception as e:
                report["failed_rows"][dataset].append(
                    {"row_number": idx, "error": str(e), "raw": {k: str(v) for k, v in row.items()}}
                )

    report["total_failed"] = sum(len(v) for v in report["failed_rows"].values())
    return report

